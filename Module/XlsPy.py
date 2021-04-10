import openpyxl as xl
import webNetPy as wp
import os,re
import ProgPy as pp
import QueryPy as qp

Erros={'xlErAbrir':100,'xlErAbaNaoEncontrada':101,'xlErPlanVazia':102,'xlErrColunas':103}

def xlLerDadosXls(caminho,abaNome):
    M=[]
    colVerify=0
    k=0
    pp.progress(50,['Carregando..'])
    while True:
        try:
            wb=xl.load_workbook(caminho,True,False)
            break
        except:
            if k==3:
                return [Erros.pop('xlErAbrir'),'Erro ao abrir a planilha de importação!\n\nArquivo não encontrado ou formato não permitido.\n\nNota: Escolha uma planilha Excel válida'],0
            k+=1

    while True:
        try:
            sh=wb.get_sheet_by_name(abaNome)
            break
        except:
            #return [Erros.pop('xlErAbaNaoEncontrada'),'Aba não encontrada na planilha de importação!\n\n "Nota: A aba deve estar nomeada como: "DADOS".'],0
            sh=wb.get_sheet_by_name(wb.get_sheet_names[0])
            
     
    i=2
    pp.progress(50,['Carregando..'])
    while sh.cell(i,1).value!=None and sh.cell(i,1).value!='':
        j=1
        M.append([])
        while sh.cell(1,j).value!=None:
            M[i-2].append(sh.cell(i,j).value)
            if j==10:
                colVerify=1
            j+=1
        i+=1
    tam=i-1
    if (tam-1)<=0:
        return [Erros.pop('xlErPlanVazia'),'A planilha importada está vazia ou não é adequada ao procedimento!'],0
    elif colVerify!=1:
        return [Erros.pop('xlErrColunas'),'A planilha importada não é adequada, ou o numero de colunas de dados a serem importados está incorreto!'],0
    else:
        return M,tam

def findLen(str): 
    counter = 0
    while str[counter:]: 
        counter += 1
    return counter

def substrRight(s,len):
    if len<findLen(s):
        aux=str(s)[(findLen(s)-len):]
    else:
        aux=s
    return aux

def substrLeft(s,len):
    if len<findLen(s):
        aux=str(s)[:len]
    else:
        aux=s
    return aux



def xlGeraVetClass(caminhoXls,abaPlan):
    M,tam=xlLerDadosXls(caminhoXls,abaPlan)

    if M[0] in [100,101,102,103]:
        return M,tam,[0]
    else:
        ret=[]
        contAnexos=[]

        for elem in M:
            urg=False
            if elem[6]==True or  str(elem[6]).upper()=='URGENTE' or str(str(elem[6]).upper()).find('URG')>-1 or str(str(elem[6]).upper()).find('SIM')>-1:
                urg=True

            paths,cont=GeraMultPaths(elem[8])
            cad=wp.kmtCadastro()
            cad.Preenche(elem[0],elem[1],elem[2],elem[3],elem[4],elem[5],urg,elem[7],paths,elem[9])

            contAnexos.append(cont)

            ret.append(cad)

        return ret,tam,contAnexos

def GeraMultPaths(caminhos):
    
    if caminhos.find(';')>=0:
        cont=str(caminhos).count(";",0)+1
        return str(caminhos).split(";"),cont
    else:
        return caminhos,1

def linhaUltima(wsInput):
    i=1
    while wsInput.cell(i,1).value!=None and wsInput.cell(i,1).value!="":
        i=i+1
    return i

def xlsExportSins(matriz,caminhoOutput="",abaOutPut=""):
   
    if caminhoOutput=="":
        wbExp=xl.Workbook()
        caminho=os.path.expanduser('~')+"\\Desktop\\SAIDAXLSKLASSMATT.xlsx"
        #wbExp.save(caminho)
        sh=wbExp.worksheets[0]
        i=1
    else:
        caminho=caminhoOutput
        wbExp=xl.load_workbook(caminhoOutput)
        if abaOutPut=="":
            sh=wbExp.get_sheet_by_name(abaOutPut)
        else:
            sh=wbExp.worksheets[0]
        i=linhaUltima(sh)
    
    for elements in matriz:
        k=1
        for var in elements:
           
            sh.cell(i,k).value=var 
            
            k=k+1
        i=i+1

    wbExp.save(caminho)


def xlExcluirLinhasCads(caminho,abaNome):
    
    k=1
    while True:
        try:
            wb=xl.load_workbook(caminho,False,False)
            break
        except:
            if k==3:
                return [Erros.pop('xlErAbrir'),'Erro ao abrir a planilha de importação!\n\nArquivo não encontrado ou formato não permitido.\n\nNota: Escolha uma planilha Excel válida'],0
            k+=1

    while True:
        try:
            sh=wb.get_sheet_by_name(abaNome)
            break
        except:
            sh=wb.get_sheet_by_name(wb.sheetnames[0])
            break

    sh.delete_rows(2,1)
    wb.save(caminho)
    wb.close()

def xlLerPlanilhaGeral(caminho,abaNome=''):
    try:
        pastatrabalho=xl.load_workbook(caminho,True)
    except:
        return

    try:
        wsFolha=pastatrabalho.get_sheet_by_name(abaNome)
    except:
        wsFolha=pastatrabalho.get_sheet_by_name(pastatrabalho.sheetnames[0])
    
    SaidaDadosPlanilha=[]

    linha=2
    coluna=1
    while wsFolha.cell(linha,1).value!=None and wsFolha.cell(linha,1).value!='' and wsFolha.cell(linha,1).value!=' ':
        coluna=1
        SaidaDadosPlanilha.append([])
        while wsFolha.cell(1,coluna).value!=None and wsFolha.cell(1,coluna).value!='' and wsFolha.cell(1,coluna).value!=' ':
            if wsFolha.cell(linha,coluna).value==None:
                SaidaDadosPlanilha[linha-2].append('0')
            else:
                SaidaDadosPlanilha[linha-2].append(wsFolha.cell(linha,coluna).value)
            coluna+=1
        linha+=1
    
    linha=linha-2
    coluna=coluna-1

    return SaidaDadosPlanilha,linha,coluna

def xlInserirDadosPlanGeral(caminho, dados, manterConteudo=True, abaNome=''):

    try:
        pastatrabalho=xl.load_workbook(caminho,False)
    except:
        return

    try:
        wsFolha=pastatrabalho.get_sheet_by_name(abaNome)
    except:
        wsFolha=pastatrabalho.get_sheet_by_name(pastatrabalho.sheetnames[0])

    if manterConteudo==False:
        if linhaUltima!=1:
            wsFolha.delete_rows(2,linhaUltima(wsFolha)-1)
        linha=2
    else:
        linha=linhaUltima(wsFolha)


    for eachElement in dados:
        coluna=1
        for eachValor in eachElement:
            wsFolha.cell(linha,coluna).value=eachValor
            coluna+=1
        linha+=1           

    #wsFolha.auto_filter.add_sort_condition("A1:G" + str(linha-1))
    pastatrabalho.save(caminho)
    pastatrabalho.close()

def obterVetColunaDeMatriz(matriz,numColuna):
    retorno=[]
    for eachElement in matriz:
        retorno.append(eachElement[numColuna])
    return retorno


def xlRemoverRegistroBaseSin(caminho,numSin, removerTudo=False):
    
    try:
        pastatrabalho=xl.load_workbook(caminho,False)
    except:
        return

    try:
        wsFolha=pastatrabalho.get_sheet_by_name(abaNome)
    except:
        wsFolha=pastatrabalho.get_sheet_by_name(pastatrabalho.sheetnames[0])

    linha=2
    while wsFolha.cell(linha,2).value!=None and wsFolha.cell(linha,2).value!='':
        
        if removerTudo==True:
            try:
                wsFolha.delete_rows(linha,1)
            except Exception as e:
                return 0,e.args[0]
            
        else:
            if wsFolha.cell(linha,2).value in numSin:
                try:
                    wsFolha.delete_rows(linha,1)
                except Exception as e:
                    return 0,e.args[0]
        linha+=1
    
    pastatrabalho.save(caminho)
    pastatrabalho.close()   

    return 1,''

def xlAtualizarRegistroBaseSin(caminho,numSin,dadosNovos):
    
    try:
        pastatrabalho=xl.load_workbook(caminho,False)
    except:
        return

    try:
        wsFolha=pastatrabalho.get_sheet_by_name(abaNome)
    except:
        wsFolha=pastatrabalho.get_sheet_by_name(pastatrabalho.sheetnames[0])

    linha=2
    while wsFolha.cell(linha,2).value!=None and wsFolha.cell(linha,2).value!='':
        if wsFolha.cell(linha,2).value in numSin:
            indice=numSin.index(wsFolha.cell(linha,2).value)
            coluna=1
            for eachElement in dadosNovos[indice]:
                wsFolha.cell(linha,coluna).value=eachElement
                coluna+=1
        linha+=1
    pastatrabalho.save(caminho)
    pastatrabalho.close()   

def ordernarPlanAcomp(caminho):

    conteudoPrevio,linhaTot,colunaTot=xlLerPlanilhaGeral(caminho)
    vetOrdenacao=[]

    for eachElement in conteudoPrevio:
        vetOrdenacao.append([eachElement[6],eachElement[0],eachElement[1],eachElement[2],eachElement[3],eachElement[4],eachElement[5]])
        
    vetOrdenacao.sort()
    conteudoNovo=[]

    linha=0
    for eachElement in vetOrdenacao:
        conteudoNovo.append([])        
        for coluna in range(1,7):
            conteudoNovo[linha].append(eachElement[coluna])
        conteudoNovo[linha].append(eachElement[0])
        linha+=1

    xlInserirDadosPlanGeral(caminho,conteudoNovo,False)
